"""Shared DOCX writing (python-docx, MIT) for PQ replies and reports."""

from __future__ import annotations

import io
import re
import time
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .config import DELIVERABLES_DIR

CITE = re.compile(r"\s*\[(\d+)\]")


def new_document() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    return doc


def centered(doc: Document, text: str, bold: bool = True, size: int = 12) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(2)


def paragraph(doc: Document, text: str, bold: bool = False, keep_citations: bool = False, italic: bool = False, size: int | None = None) -> None:
    for chunk in (text or "").split("\n"):
        line = chunk if keep_citations else CITE.sub("", chunk)
        line = line.replace("**", "")
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if line.lstrip().startswith("- "):
            line = "• " + line.lstrip()[2:]
        run = p.add_run(line)
        run.bold = bold
        run.italic = italic
        if size:
            run.font.size = Pt(size)


def table(doc: Document, header: list[str], rows: list[list[str]], title: str | None = None) -> None:
    if title:
        centered(doc, title, bold=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(str(text))
        run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
    doc.add_paragraph()


def image(doc: Document, png: bytes, width_inches: float = 6.0) -> None:
    from docx.shared import Inches

    doc.add_picture(io.BytesIO(png), width=Inches(width_inches))


def evidence_section(doc: Document, citations: list[dict], heading: str = "Evidence trail (internal — remove before dispatch)") -> None:
    if not citations:
        return
    doc.add_page_break()
    p = doc.add_paragraph()
    run = p.add_run(heading)
    run.bold = True
    run.font.color.rgb = RGBColor(0x80, 0x40, 0x00)
    for c in citations:
        where = f"page {c['page_no']}" if c.get("page_no") else ""
        status = f" · {c['status']}" if c.get("status") else ""
        paragraph(doc, f"[{c['n']}] {c['filename']} {where}{status} — {c.get('snippet', '')[:300]}", keep_citations=True, size=9)


def save(doc: Document, stem: str, filename: str | None = None) -> Path:
    """Write the document to deliverables/. A caller-chosen `filename` (the harness tool picks one up front so
    its deliverables row can point at the file before it exists) replaces the timestamped default."""
    DELIVERABLES_DIR.mkdir(parents=True, exist_ok=True)
    if filename:
        name = re.sub(r"[^A-Za-z0-9._-]+", "-", Path(filename).name).strip("-")
        name = name if name.lower().endswith(".docx") else f"{name}.docx"
    else:
        safe = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-")[:80] or "stratum"
        name = f"{safe}-{time.strftime('%Y%m%d-%H%M%S')}.docx"
    path = DELIVERABLES_DIR / name
    doc.save(path)
    return path


def table_rows(t: dict) -> tuple[list[str], list[list[str]]]:
    """An Ask answer table → header + rows for DOCX."""
    achievement = t.get("kind") == "achievement"
    header = ["Entity", *t["periods"]] + ([t.get("change_label", "Change %").replace("%", "(%)")] if t.get("show_change") else [])
    rows = []
    for row in t["rows"]:
        cells = [(c["display"] + (" (P)" if c["provisional"] else "")) if c else "—" for c in row["cells"]]
        pct = row.get("change_pct")
        extra = [("—" if pct is None else (f"{pct:.2f}" if achievement else f"{pct:+.2f}"))] if t.get("show_change") else []
        rows.append([row["entity"], *cells, *extra])
    return header, rows
